from celery import shared_task
from .models import CVUpload, CompetenceDossier
from pyresparser import ResumeParser
from docx import Document
import os

@shared_task
def generate_dossier(cv_id):
    cv = CVUpload.objects.get(id=cv_id)
    data = ResumeParser(cv.file.path).get_extracted_data()

    # Génération du Word (simplifié)
    doc = Document()
    doc.add_heading(f"Dossier de {data['name']}", level=1)
    doc.add_paragraph("Compétences: " + ", ".join(data['skills']))
    output_path = f"media/dossiers/dossier_{cv_id}.docx"
    doc.save(output_path)

    # Sauvegarde dans la base
    CompetenceDossier.objects.create(
        cv=cv,
        generated_file=output_path.replace("media/", "")
    )
